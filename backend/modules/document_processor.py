"""
Document Processing Module

This module handles document text extraction and processing for the DocuMind application.
Supports PDF, DOCX, TXT files with pluggable architecture for easy testing and extension.

Architecture:
- DocumentProcessor: Main orchestrator class
- DocumentLoader: File format specific text extraction
- TextChunker: Semantic text splitting for RAG
"""

from typing import List, Dict, Any, Optional
from pathlib import Path
import mimetypes
from abc import ABC, abstractmethod


class DocumentLoader(ABC):
    """Abstract base class for document loaders"""
    
    @abstractmethod
    def can_process(self, file_path: str, mime_type: str) -> bool:
        """Check if this loader can process the given file type"""
        pass
    
    @abstractmethod
    def extract_text(self, file_path: str) -> str:
        """Extract text from the document"""
        pass
    
    @abstractmethod
    def extract_metadata(self, file_path: str) -> Dict[str, Any]:
        """Extract metadata from the document"""
        pass


class PlainTextLoader(DocumentLoader):
    """Loader for plain text files (TXT, MD)"""
    
    def can_process(self, file_path: str, mime_type: str) -> bool:
        return mime_type.startswith('text/') or file_path.lower().endswith(('.txt', '.md'))
    
    def extract_text(self, file_path: str) -> str:
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except UnicodeDecodeError:
            # Fallback to latin-1 encoding
            with open(file_path, 'r', encoding='latin-1') as f:
                return f.read()
    
    def extract_metadata(self, file_path: str) -> Dict[str, Any]:
        path = Path(file_path)
        return {
            "format": "text",
            "pages": 1,
            "characters": len(self.extract_text(file_path)),
            "file_size": path.stat().st_size
        }


class PDFLoader(DocumentLoader):
    """Loader for PDF files using PyPDF2 or pdfplumber"""
    
    def can_process(self, file_path: str, mime_type: str) -> bool:
        return mime_type == 'application/pdf' or file_path.lower().endswith('.pdf')
    
    def extract_text(self, file_path: str) -> str:
        try:
            # Try with pdfplumber first (better text extraction)
            import pdfplumber
            text = ""
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n\n"
            return text.strip()
        except ImportError:
            # Fallback to PyPDF2
            try:
                import PyPDF2
                text = ""
                with open(file_path, 'rb') as file:
                    reader = PyPDF2.PdfReader(file)
                    for page in reader.pages:
                        text += page.extract_text() + "\n\n"
                return text.strip()
            except ImportError:
                # If no PDF libraries available, return placeholder
                return f"[PDF text extraction unavailable - install pdfplumber or PyPDF2]\nFile: {file_path}"
    
    def extract_metadata(self, file_path: str) -> Dict[str, Any]:
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                return {
                    "format": "pdf",
                    "pages": len(pdf.pages),
                    "characters": len(self.extract_text(file_path)),
                    "file_size": Path(file_path).stat().st_size
                }
        except ImportError:
            try:
                import PyPDF2
                with open(file_path, 'rb') as file:
                    reader = PyPDF2.PdfReader(file)
                    return {
                        "format": "pdf",
                        "pages": len(reader.pages),
                        "characters": len(self.extract_text(file_path)),
                        "file_size": Path(file_path).stat().st_size
                    }
            except ImportError:
                return {
                    "format": "pdf",
                    "pages": 0,
                    "characters": 0,
                    "file_size": Path(file_path).stat().st_size,
                    "error": "No PDF processing library available"
                }


class DOCXLoader(DocumentLoader):
    """Loader for DOCX files using python-docx"""
    
    def can_process(self, file_path: str, mime_type: str) -> bool:
        return (mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' 
                or file_path.lower().endswith('.docx'))
    
    def extract_text(self, file_path: str) -> str:
        try:
            from docx import Document
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text.strip()
        except ImportError:
            return f"[DOCX text extraction unavailable - install python-docx]\nFile: {file_path}"
    
    def extract_metadata(self, file_path: str) -> Dict[str, Any]:
        try:
            from docx import Document
            doc = Document(file_path)
            text = self.extract_text(file_path)
            return {
                "format": "docx",
                "paragraphs": len(doc.paragraphs),
                "characters": len(text),
                "file_size": Path(file_path).stat().st_size
            }
        except ImportError:
            return {
                "format": "docx",
                "paragraphs": 0,
                "characters": 0,
                "file_size": Path(file_path).stat().st_size,
                "error": "python-docx library not available"
            }


from modules.rag_pipeline import TextChunker
from sentence_transformers import SentenceTransformer
import numpy as np

class DocumentProcessor:
    """Main document processing orchestrator (now also handles chunking and embedding)"""

    def __init__(self):
        self.loaders = [
            PlainTextLoader(),
            PDFLoader(),
            DOCXLoader()
        ]
        # Use same chunking and embedding config as RAGPipeline default
        self.chunker = TextChunker(chunk_size=1000, overlap=200)
        self.embedding_model = SentenceTransformer('all-MiniLM-L6-v2')

    def get_loader_for_file(self, file_path: str) -> Optional[DocumentLoader]:
        """Find appropriate loader for the file type"""
        mime_type, _ = mimetypes.guess_type(file_path)
        mime_type = mime_type or ""

        for loader in self.loaders:
            if loader.can_process(file_path, mime_type):
                return loader
        return None

    def process_document(self, file_path: str) -> Dict[str, Any]:
        """
        Process a document: extract text, metadata, chunk, and embed at upload time.

        Returns:
            Dict containing:
            - text: Full extracted text
            - metadata: Document metadata
            - status: Processing status
            - error: Error message if processing failed
            - chunks: List of chunk dicts (with text, length, metadata)
            - embeddings: np.ndarray of chunk embeddings (n_chunks x embedding_dim)
        """
        try:
            loader = self.get_loader_for_file(file_path)
            if not loader:
                return {
                    "text": "",
                    "metadata": {},
                    "status": "unsupported",
                    "error": f"Unsupported file type: {Path(file_path).suffix}",
                    "chunks": [],
                    "embeddings": None
                }
            text = loader.extract_text(file_path)
            metadata = loader.extract_metadata(file_path)
            # Chunk and embed here
            if text.strip():
                chunks = self.chunker.chunk_text(text, metadata)
                if chunks:
                    embeddings = self.embedding_model.encode([c['text'] for c in chunks], convert_to_numpy=True)
                else:
                    embeddings = np.empty((0, self.embedding_model.get_sentence_embedding_dimension()), dtype=np.float32)
            else:
                chunks = []
                embeddings = np.empty((0, self.embedding_model.get_sentence_embedding_dimension()), dtype=np.float32)
            return {
                "text": text,
                "metadata": metadata,
                "status": "processed",
                "error": None,
                "chunks": chunks,
                "embeddings": embeddings
            }
        except Exception as e:
            return {
                "text": "",
                "metadata": {},
                "status": "error",
                "error": str(e),
                "chunks": [],
                "embeddings": None
            }

    def get_supported_formats(self) -> List[str]:
        """Return list of supported file formats"""
        return ['.pdf', '.docx', '.txt', '.md']

    def is_supported(self, file_path: str) -> bool:
        """Check if file format is supported"""
        return self.get_loader_for_file(file_path) is not None